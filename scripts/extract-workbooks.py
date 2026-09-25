"""Stream source workbooks to JSONL for an admin-only database import.

Usage: python scripts/extract-workbooks.py /path/to/excels > /tmp/rerchar-workbooks.jsonl
Requires: pip install openpyxl. The output contains sensitive customer data; keep it private.
"""
import datetime
import hashlib
import json
import math
import pathlib
import sys

from openpyxl import load_workbook


def normalized(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (int, float, bool, str)) or value is None:
        return value
    return str(value)


def emit(value):
    print(json.dumps(value, ensure_ascii=False, separators=(",", ":")), flush=True)


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Uso: extract-workbooks.py DIRECTORIO_DE_EXCEL")
    directory = pathlib.Path(sys.argv[1]).resolve()
    if not directory.is_dir():
        raise SystemExit(f"No existe el directorio: {directory}")
    paths = sorted(p for p in directory.iterdir() if p.suffix.lower() in (".xlsx", ".xlsm") and not p.name.startswith("~$"))
    if not paths:
        raise SystemExit("No hay libros .xlsx ni .xlsm en el directorio")
    for path in paths:
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        emit({"type": "source", "filename": path.name, "sha256": digest})
        formulas = load_workbook(path, read_only=True, data_only=False)
        cached = load_workbook(path, read_only=True, data_only=True)
        for formula_sheet, cached_sheet in zip(formulas.worksheets, cached.worksheets):
            if formula_sheet.title != cached_sheet.title:
                raise RuntimeError("Hojas desalineadas: " + path.name)
            for row_number, (formula_row, cached_row) in enumerate(zip(formula_sheet.iter_rows(), cached_sheet.iter_rows()), start=1):
                cells = []
                for fc, cc in zip(formula_row, cached_row):
                    if fc.value is None and cc.value is None:
                        continue
                    value = normalized(cc.value if cc.value is not None else fc.value)
                    if value == "" or value is None:
                        continue
                    cell = {"col": fc.column_letter, "value": value}
                    if fc.data_type == "f":
                        cell["formula"] = str(fc.value)
                    cells.append(cell)
                if cells:
                    emit({"type": "row", "sheet": formula_sheet.title, "row": row_number, "cells": cells})
        cached.close()
        formulas.close()
        emit({"type": "end"})


if __name__ == "__main__":
    main()
